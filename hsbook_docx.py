# -*- coding: utf-8 -*-
"""Reusable toolkit for building Chinese-translation .docx files for
'Hydraulic Servo-systems' (Jelali & Kroll, 2003).

中文为主 + 关键术语中英对照。Shared by all per-chapter build scripts.
"""
import os
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LATIN = "Times New Roman"
CJK = "宋体"
CJK_BOLD = "黑体"

SRC_PDF = ("(Advances in Industrial Control) Mohieddine Jelali Dr-Ing, "
           "Andreas Kroll Dr-Ing (auth.) - Hydraulic Servo-systems_ Modelling, "
           "Identification and Control-Springer-Verlag London (2003).pdf")

# printed page N  <->  PDF page index N + 26
def pidx(printed_page):
    return printed_page + 26


def crop_region(pdf_path, page_index, ytop, ybot, xleft, xright, out_path, dpi=200):
    """Crop a fractional region of a PDF page to a PNG."""
    doc = fitz.open(pdf_path)
    pg = doc[page_index]
    r = pg.rect
    clip = fitz.Rect(r.x0 + xleft * r.width, r.y0 + ytop * r.height,
                     r.x0 + xright * r.width, r.y0 + ybot * r.height)
    pg.get_pixmap(dpi=dpi, clip=clip).save(out_path)
    doc.close()
    return out_path


class DocBuilder:
    def __init__(self):
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_height = Cm(29.7)
        sec.page_width = Cm(21.0)
        for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, m, Cm(2.4))
        self.content_w = sec.page_width - sec.left_margin - sec.right_margin

        normal = self.doc.styles["Normal"]
        normal.font.name = LATIN
        normal.font.size = Pt(12)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        normal.paragraph_format.space_after = Pt(0)
        for sid, sz in (("Heading 1", 15), ("Heading 2", 13.5), ("Heading 3", 12)):
            st = self.doc.styles[sid]
            st.font.name = "Arial"; st.font.size = Pt(sz); st.font.bold = True
            st.font.color.rgb = RGBColor(0, 0, 0)
            rpr = st.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.append(rf)
            rf.set(qn("w:eastAsia"), CJK_BOLD)

    # ---- run-level font ----
    def _set_run(self, run, latin=LATIN, cjk=CJK, size=None, bold=None,
                 italic=None, color=None):
        run.font.name = latin
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        rf.set(qn("w:ascii"), latin); rf.set(qn("w:hAnsi"), latin)
        rf.set(qn("w:eastAsia"), cjk)
        if size is not None: run.font.size = Pt(size)
        if bold is not None: run.font.bold = bold
        if italic is not None: run.font.italic = italic
        if color is not None: run.font.color.rgb = color

    # ---- math markup: _{...} subscript, ^{...} superscript ----
    def _add_math(self, p, markup, size=12, color=None, italic=False):
        i, n, buf = 0, len(markup), ""
        def flush():
            nonlocal buf
            if buf:
                self._set_run(p.add_run(buf), size=size, color=color, italic=italic)
                buf = ""
        while i < n:
            ch = markup[i]
            if ch in "_^" and i + 1 < n and markup[i + 1] == "{":
                flush()
                j = markup.index("}", i + 2)
                r = p.add_run(markup[i + 2:j])
                self._set_run(r, size=size, color=color, italic=italic)
                if ch == "_": r.font.subscript = True
                else: r.font.superscript = True
                i = j + 1
            else:
                buf += ch; i += 1
        flush()

    # ---- blocks ----
    def title(self, cn, en, meta_lines):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        self._set_run(p.add_run(cn), latin="Arial", cjk=CJK_BOLD, size=19, bold=True)
        p2 = self.doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._set_run(p2.add_run(en), size=11, italic=True, color=RGBColor(0x44,0x44,0x44))
        for line in meta_lines:
            pl = self.doc.add_paragraph(); pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pl.paragraph_format.space_after = Pt(2)
            self._set_run(pl.add_run(line), size=10.5, color=RGBColor(0x33,0x33,0x33))

    def note_box(self, title_cn, body):
        self.doc.add_paragraph()
        box = self.doc.add_paragraph()
        box.paragraph_format.space_before = Pt(6)
        pPr = box._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        for edge in ("top", "left", "bottom", "right"):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
            e.set(qn("w:space"), "6"); e.set(qn("w:color"), "999999")
            pbdr.append(e)
        pPr.insert(0, pbdr)
        self._set_run(box.add_run(title_cn), bold=True, size=10.5)
        self._set_run(box.add_run("\n" + body), size=10.5, color=RGBColor(0x33,0x33,0x33))

    def h1(self, cn, en=None):
        p = self.doc.add_heading(level=1)
        self._set_run(p.add_run(cn), latin="Arial", cjk=CJK_BOLD, size=15, bold=True,
                      color=RGBColor(0,0,0))
        if en:
            p.add_run("  ")
            self._set_run(p.add_run(en), latin="Arial", cjk=CJK_BOLD, size=10.5,
                          bold=False, color=RGBColor(0x55,0x55,0x55))

    def h2(self, cn, en=None):
        p = self.doc.add_heading(level=2)
        self._set_run(p.add_run(cn), latin="Arial", cjk=CJK_BOLD, size=13.5, bold=True,
                      color=RGBColor(0,0,0))
        if en:
            p.add_run("  ")
            self._set_run(p.add_run(en), latin="Arial", cjk=CJK_BOLD, size=10,
                          bold=False, color=RGBColor(0x55,0x55,0x55))

    def h3(self, cn, en=None):
        p = self.doc.add_heading(level=3)
        self._set_run(p.add_run(cn), latin="Arial", cjk=CJK_BOLD, size=12, bold=True,
                      color=RGBColor(0,0,0))
        if en:
            p.add_run("  ")
            self._set_run(p.add_run(en), latin="Arial", cjk=CJK_BOLD, size=9.5,
                          bold=False, color=RGBColor(0x55,0x55,0x55))

    def para(self, cn, indent=True):
        p = self.doc.add_paragraph()
        if indent:
            p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.space_after = Pt(4)
        self._set_run(p.add_run(cn))
        return p

    def para_rich(self, segments, indent=True):
        """segments: list of (text, kind) kind in {'',' b','i','m'} m=math markup"""
        p = self.doc.add_paragraph()
        if indent:
            p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.space_after = Pt(4)
        for text, kind in segments:
            if kind == "m":
                self._add_math(p, text)
            else:
                self._set_run(p.add_run(text), bold=(kind == "b"), italic=(kind == "i"))
        return p

    def bullet(self, cn):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        self._set_run(p.add_run(cn))
        return p

    def label(self, cn):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(0)
        self._set_run(p.add_run(cn), bold=True)
        return p

    def equation(self, markup, number=None):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(6); pf.space_after = Pt(6)
        pf.tab_stops.add_tab_stop(int(self.content_w / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(int(self.content_w), WD_TAB_ALIGNMENT.RIGHT)
        self._set_run(p.add_run("\t"))
        self._add_math(p, markup)
        if number:
            self._set_run(p.add_run("\t(" + number + ")"))
        return p

    def note(self, markup):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        self._add_math(p, markup, size=11, color=RGBColor(0x33,0x33,0x33))

    def figure(self, img_path, num, cn, en, width_cm):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(img_path, width=Cm(width_cm))
        cap = self.doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        self._set_run(cap.add_run(f"图 {num}　{cn}"), size=10.5, bold=True)
        self._set_run(cap.add_run(f"\n(Figure {num}. {en})"), size=9, italic=True,
                      color=RGBColor(0x66,0x66,0x66))

    def image_block(self, img_path, width_cm, caption_cn=None):
        """Embed an image (e.g., an equation/table cropped from the scan)."""
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(img_path, width=Cm(width_cm))
        if caption_cn:
            cap = self.doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._set_run(cap.add_run(caption_cn), size=9.5, color=RGBColor(0x55,0x55,0x55))

    def _shade(self, cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

    def table(self, headers, rows, col_widths_cm, math_cols=()):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, txt in zip(t.rows[0].cells, headers):
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._set_run(c.paragraphs[0].add_run(txt), size=10.5, bold=True)
            self._shade(c, "D9E2F3")
        for row in rows:
            cells = t.add_row().cells
            for k, val in enumerate(row):
                par = cells[k].paragraphs[0]
                if k in math_cols:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self._add_math(par, val, size=10)
                else:
                    self._set_run(par.add_run(val), size=10)
        for row in t.rows:
            for k, c in enumerate(row.cells):
                c.width = Cm(col_widths_cm[k])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        # schema fix: ensure zoom has percent
        s = self.doc.settings.element
        z = s.find(qn("w:zoom"))
        if z is not None and z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
        self.doc.save(path)
        return path
