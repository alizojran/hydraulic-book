#!/usr/bin/env python3
"""Translation workflow for Servohydraulik.

Subcommands:
  prep            build batches.json (ordered key chunks ~TARGET chars each)
  show N          print German text of batch N (untranslated items)
  status          progress report
  assemble [out]  build the Chinese .docx from extracted.json + translations/
"""
import json, os, sys, glob, re

DIR = os.path.dirname(os.path.abspath(__file__))
EXTRACTED = os.path.join(DIR, "extracted.json")
BATCHES = os.path.join(DIR, "batches.json")
TRDIR = os.path.join(DIR, "translations")
TARGET = 16000   # German chars per batch
os.makedirs(TRDIR, exist_ok=True)


def load_extracted():
    with open(EXTRACTED, encoding="utf-8") as f:
        return json.load(f)


def ordered_text_elements(data):
    """Flat list of (k, type, german) in document order."""
    out = []
    for p in data["pages"]:
        for e in p["elements"]:
            if "k" in e:
                out.append((e["k"], e["t"], e["de"]))
    return out


def load_translations():
    """Merge every translations/*.json into one {int_key: chinese}."""
    tr = {}
    for fp in sorted(glob.glob(os.path.join(TRDIR, "*.json"))):
        with open(fp, encoding="utf-8") as f:
            d = json.load(f)
        for k, v in d.items():
            tr[int(k)] = v
    return tr


def cmd_prep():
    data = load_extracted()
    elems = ordered_text_elements(data)
    batches, cur, size = [], [], 0
    for k, t, de in elems:
        cur.append(k)
        size += len(de)
        # headings start fresh-ish but keep simple: flush on size
        if size >= TARGET:
            batches.append(cur); cur = []; size = 0
    if cur:
        batches.append(cur)
    with open(BATCHES, "w", encoding="utf-8") as f:
        json.dump({"batches": batches}, f)
    print(f"{len(elems)} text elements -> {len(batches)} batches (~{TARGET} chars each)")
    # per-batch char sizes
    by_key = {k: de for k, t, de in elems}
    for i, b in enumerate(batches):
        print(f"  batch {i:02d}: {len(b):3d} items, {sum(len(by_key[k]) for k in b):5d} chars")


def cmd_show(n):
    data = load_extracted()
    by_key = {k: (t, de) for k, t, de in ordered_text_elements(data)}
    with open(BATCHES, encoding="utf-8") as f:
        batches = json.load(f)["batches"]
    tr = load_translations()
    b = batches[n]
    todo = [k for k in b if k not in tr]
    print(f"# BATCH {n} / {len(batches)-1}  ({len(todo)} of {len(b)} untranslated)")
    print(f"# write translations to: translations/t_{n:03d}.json  as {{\"<k>\": \"中文\", ...}}")
    print("#" + "=" * 70)
    for k in b:
        if k in tr:
            continue
        t, de = by_key[k]
        print(f"[{k}|{t}] {de}")


def cmd_status():
    data = load_extracted()
    elems = ordered_text_elements(data)
    tr = load_translations()
    done = sum(1 for k, t, de in elems if k in tr)
    done_chars = sum(len(de) for k, t, de in elems if k in tr)
    tot_chars = sum(len(de) for k, t, de in elems)
    with open(BATCHES, encoding="utf-8") as f:
        batches = json.load(f)["batches"]
    pending = [i for i, b in enumerate(batches) if any(k not in tr for k in b)]
    print(f"elements : {done}/{len(elems)} translated ({100*done/len(elems):.1f}%)")
    print(f"chars    : {done_chars}/{tot_chars} ({100*done_chars/tot_chars:.1f}%)")
    print(f"batches  : {len(batches)-len(pending)}/{len(batches)} done")
    if pending:
        print(f"pending batches: {pending}")


# ---------------- docx assembly ----------------
def set_cjk_font(doc):
    from docx.oxml.ns import qn
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3",
                       "Title", "Subtitle"):
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            from docx.oxml import OxmlElement
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        rf.set(qn("w:eastAsia"), "Microsoft YaHei")
        rf.set(qn("w:ascii"), "Microsoft YaHei")
        rf.set(qn("w:hAnsi"), "Microsoft YaHei")


def cmd_assemble(out):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    data = load_extracted()
    tr = load_translations()
    doc = docx.Document()
    set_cjk_font(doc)
    # base font size
    doc.styles["Normal"].font.size = Pt(10.5)

    san = lambda s: re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', ' ', s)
    HSTYLE = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}
    MAXW = 6.0  # inches
    missing = 0
    for p in data["pages"]:
        for e in p["elements"]:
            if e["t"] == "fig":
                path = os.path.join(DIR, e["img"])
                if not os.path.exists(path):
                    continue
                wn = e.get("w", 200) / 72.0
                wn = min(wn, MAXW)
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    par.add_run().add_picture(path, width=Inches(wn))
                except Exception:
                    pass
                continue
            k = e["k"]
            zh = tr.get(k)
            if zh is None:
                missing += 1
                zh = "〔待译〕" + e["de"]
            zh = san(zh).strip()
            if not zh:
                continue
            t = e["t"]
            if t in HSTYLE:
                doc.add_paragraph(zh, style=HSTYLE[t])
            elif t == "cap":
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = par.add_run(zh)
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            else:
                doc.add_paragraph(zh)
    doc.save(out)
    print(f"saved {out}  ({missing} untranslated text elements fell back to German)")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "status"
    if cmd == "prep":
        cmd_prep()
    elif cmd == "show":
        cmd_show(int(sys.argv[2]))
    elif cmd == "status":
        cmd_status()
    elif cmd == "assemble":
        out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
            os.path.dirname(DIR), "Servohydraulik_中文翻译.docx")
        cmd_assemble(out)
    else:
        print("unknown command", cmd)
