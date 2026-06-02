# -*- coding: utf-8 -*-
"""Front matter: title page, translated preface, bilingual table of contents."""
from viersma_docx import BookBuilder
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

b = BookBuilder()

# ---------------- Title page ----------------
b.title(
    "液压伺服系统与管路的分析、综合与设计",
    "Analysis, Synthesis and Design of Hydraulic Servosystems and Pipelines",
    [
        "Taco J. Viersma",
        "（机械工程研究丛书 第 1 卷 · Studies in Mechanical Engineering 1）",
        "Elsevier Scientific Publishing Company，阿姆斯特丹，1980 年",
        "中文译本",
    ],
)

note = b.doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
note.paragraph_format.space_before = Pt(24)
b._set_run(note.add_run(
    "说明：原书为扫描影印件（无文字层）。本译本中正文、标题与图题均译为中文，"
    "关键术语保留中英对照；公式依据原书重新排版（数学符号沿用国际通用写法），"
    "插图自原书扫描页裁切嵌入并附中文图题。"),
    size=10.5, italic=True, color=RGBColor(0x55, 0x55, 0x55))

b.page_break()

# ---------------- Preface ----------------
b.h1("前言", "Preface")

for para in [
    "本书是代尔夫特理工大学（Delft University of Technology）一门新开设的液压伺服系统"
    "研究生课程的总结。之所以以英文写成，是因为许多国家的科学家、学生、工程师和制造商，"
    "都对这一主题怀有切实而迫切的兴趣。",

    "本书内容的主要来源有：",
]:
    b.para(para)

b.bullet("代尔夫特理工大学机械工程系机械控制工程实验室液压研究组的研究与开发成果；")
b.bullet("T. J. Viersma 的博士论文《液压伺服马达精度研究》"
         "（Investigations into the accuracy of hydraulic servomotors），"
         "代尔夫特理工大学，1961 年；")
b.bullet("P. Blok 的博士论文《锥形静压轴承》"
         "（Conical hydrostatic bearings，德文），代尔夫特理工大学，1976 年；")
b.bullet("A. A. Ham 关于《液压供油管路动力学》"
         "（The dynamics of hydraulic supply lines）的博士论文（撰写中）。")

for para in [
    "本书所讨论问题的一个重要基础，是 Blackburn、Lee 和 Shearer 在其杰作"
    "《流体动力控制》（Fluid Power Control，M.I.T. Technology Press，1960）中奠定的。"
    "这部七百余页的著作虽已无法就本主题提供完整而最新的信息，但它仍是一座几乎"
    "取之不尽的基础知识宝库。",

    "其次值得推荐的，是 Merritt 的《液压控制系统》"
    "（Hydraulic Control Systems，Wiley，纽约，1967）。这部 350 页的著作更侧重于"
    "液压控制问题。除最后几章外，全书都极为出色；但它并未触及设计阶段。",

    "在本书中，我给自己定下的目标是：用极其简短、紧凑的篇幅，总结我所掌握的有关"
    "液压伺服系统分析、综合与设计的知识。依托代尔夫特理工大学液压研究组的专业技术"
    "与经验，我尽可能直截了当地、沿最短的路线奔向"
    "“设计”这一目标。",

    "在伺服系统的基础内容之外，本书还增加了两个专门的章节。第 6 章讨论伺服系统"
    "液压供油管路的动力学；文献中一向假定供油压力恒定，却从不说明如何实现这一点。"
    "第 7 章总结了 Blok 在静压轴承方面的杰出工作。静压轴承几乎可以完全消除库仑摩擦，"
    "从而避免死区或黏滑（slip-stick）运动等现象。由此带来的、对代尔夫特与阿姆斯特丹"
    "飞行模拟器性能的显著改善，促使许多外国制造商去满足若干欧洲主要航空公司的高标准要求。",

    "自学第 1～5 章的读者，可以略去所有带星号（＊）的小节。这些小节涉及的细节，"
    "对于理解其余内容并非必不可少。本书中的插图具有特殊作用：在通读全书、对内容有了"
    "透彻了解之后，左页上九十余幅图几乎可以作为全书内容的完整图解纲要。",

    "最后，我谨向直接给予支持的同事 Teerhuis 和 Reuvekamp，向负责绘图的 Gutteling "
    "和 Jansen，以及负责录入文稿的 Bosch 女士，致以诚挚的谢意。",
]:
    b.para(para)

sign = b.doc.add_paragraph()
sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
sign.paragraph_format.space_before = Pt(12)
b._set_run(sign.add_run("代尔夫特，1979 年 12 月\nTaco J. Viersma"),
           italic=True)

b.page_break()

# ---------------- Table of contents ----------------
b.h1("目录", "Contents")

CH = lambda n, cn, en: b.para_rich(
    [(f"第 {n} 章　", "b"), (cn + "  ", "b"), (en, "i")], indent=False)


def SEC(num, cn, en, star=False):
    pre = "＊" if star else ""
    p = b.para_rich([(f"　　{pre}{num}　", ""), (cn + "  ", ""),
                     ("（" + en + "）", "i")], indent=False)
    p.paragraph_format.space_after = Pt(1)
    return p


CH(1, "引论", "Introduction")
SEC("1.1", "带比例反馈的积分器", "Integrator with proportional feedback")
SEC("1.2", "简单伺服系统的主要特性", "Dominant properties of a simple servosystem")
SEC("1.3", "速度增益 Kᵥ 与阀的配置", "Velocity gain Kv and valve configuration")
SEC("1.4", "多孔阀", "Multiple-hole valve", star=True)

CH(2, "液压伺服系统的分析工具", "Tools for the analysis of hydraulic servo's")
SEC("2.1", "湍流口流量", "Turbulent port flow")
SEC("2.2", "油液的可压缩性", "Compressibility of oil")
SEC("2.3", "力与连续性", "Forces and continuity")

CH(3, "液压伺服系统的分析", "Analysis of hydraulic servosystems")
SEC("3.1", "引言", "Introduction")
SEC("3.2", "三通阀控制的非对称马达", "Asymmetric motors controlled by 3-way valves")
SEC("3.3", "负载补偿", "Load-compensation", star=True)
SEC("3.4", "对称四通阀控制的对称马达",
    "Symmetric motor controlled by a symmetric 4-way valve")
SEC("3.5", "对称四通零开口阀控制的非对称马达",
    "Asymmetric motor controlled by a symmetric 4-way critical-center valve", star=True)
SEC("3.6", "非对称四通阀控制的非对称马达",
    "Asymmetric motor controlled by an asymmetric 4-way valve", star=True)
SEC("3.7", "从数学模型到传递函数", "From mathematical model to transfer function")

CH(4, "动态分析与综合", "Dynamic analysis and synthesis")
SEC("4.1", "引言", "Introduction")
SEC("4.2", "比例控制", "Proportional control")
SEC("4.3", "动态负载", "Dynamic load")
SEC("4.4", "库仑摩擦引起的动态死区", "Dynamic dead zone due to Coulomb friction", star=True)
SEC("4.5", "阻尼方法", "Damping methods")
SEC("4.6", "电液伺服阀", "Electro-hydraulic servovalve")
SEC("4.7", "硬件元件的方框图", "Block diagram for hardware components")
SEC("4.8", "PI 控制", "PI-control", star=True)
SEC("4.9", "串联一阶系统", "First order system in series", star=True)
SEC("4.10", "柔顺结构", "Compliant structures", star=True)
SEC("4.11", "复杂液压伺服系统的低频控制", "LF-control of complex hydraulic servo's", star=True)
SEC("4.12", "负载作动器", "Load actuators", star=True)
SEC("4.13", "库仑摩擦与稳定性", "Coulomb friction and stability", star=True)
SEC("4.14", "库仑摩擦的模拟仿真", "Analog simulation of Coulomb friction", star=True)
SEC("4.15", "库仑摩擦的描述函数", "Describing function of Coulomb friction", star=True)

CH(5, "受惯性力与恒定预载作用的正弦驱动线性伺服系统——性能图与设计",
   "Performance diagram and design of sinusoidally actuated linear servosystems "
   "loaded by inertia forces and a constant pre-load")
SEC("5.1", "引言", "Introduction")
SEC("5.2", "最大行程与带宽的关系", "Relation between max. stroke and bandwidth")
SEC("5.3", "阀流量能力与最大速度的关系",
    "Relation between valve capacity and maximum speed")
SEC("5.4", "气穴边界与最大加速度的关系",
    "Relation between cavitation boundary and max. acceleration")
SEC("5.5", "初步性能图", "Preliminary performance diagram")
SEC("5.6", "对称马达的设计", "Design of symmetric motors")
SEC("5.7", "最终性能图", "Final performance diagram")
SEC("5.8", "非对称伺服马达的设计", "Design of asymmetric servomotors", star=True)
SEC("5.9", "设计算例", "Design example")
SEC("5.10", "反向设计算例", "Inverse design example")

CH(6, "液压管路动力学", "Hydraulic line dynamics")
SEC("6.1", "引言", "Introduction")
SEC("6.2", "液压管路动态行为的基本方程",
    "Basic equations for dynamic behavior of hydraulic lines")
SEC("6.3", "单根管路的动态行为", "Dynamic behavior of a single pipeline")
SEC("6.4", "复合管路与元件的一般规则",
    "General rules for compound pipelines and components")
SEC("6.5", "泵流量脉动的抑制", "Suppression of pump flow pulsations")
SEC("6.6", "减小伺服阀耗流引起的供油压力波动",
    "Reducing supply pressure variations due to consumption by the servovalve")
SEC("6.7", "仿真", "Simulation")

CH(7, "高速液压缸用锥形静压轴承的设计",
   "The design of conical hydrostatic bearings for hydraulic cylinders "
   "with considerable speed")
SEC("7.1", "引言", "Introduction")
SEC("7.2", "单轴承", "Single bearings")
SEC("7.3", "推动应用的新设计突破", "New designs as a breakthrough to applications")
SEC("7.4", "双轴承", "Double bearings")
SEC("7.5", "制造", "Manufacturing")
SEC("7.6", "应用", "Applications")
SEC("7.7", "结论与小结", "Conclusions and summary")
SEC("7.8", "附录", "Appendix")
SEC("7.9", "设计算例", "Design example")
SEC("7.10", "雷诺方程的推导", "Derivation of Reynolds' equation", star=True)
SEC("7.11", "压力分布的推导", "Derivation of pressure distribution", star=True)
SEC("7.12", "轴承力的推导", "Derivation of bearing force", star=True)
SEC("7.13", "泄漏流量的推导", "Derivation of leakage flow", star=True)
SEC("7.14", "双轴承气穴边界的推导",
    "Derivation of cavitation boundary for double bearings", star=True)
SEC("7.15", "锥形轴承二维流动的迭代松弛解法",
    "Iterative relaxation for conical bearing with two-dimensional flow", star=True)
SEC("7.16", "二维流动计算结果", "Results of calculations on two-dimensional flow", star=True)

tip = b.doc.add_paragraph()
tip.paragraph_format.space_before = Pt(10)
b._set_run(tip.add_run("注：标有“＊”的小节为深入内容，自学第 1～5 章时可略去而不影响理解。"),
           size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66))

b.save("parts/front.docx")
print("saved parts/front.docx")
